"""Controles sintéticos locales. Ninguna API, cuenta o documento del usuario."""
import importlib.util
import io
import json
import os
from pathlib import Path
import subprocess
import sys
from unittest import TestCase
from unittest.mock import Mock, patch
import zipfile

from conversion import local_engines as le
from conversion.documents import digest, json_bytes, load_json
from conversion.storage import verify_artifacts
from tests.support import workspace_temp


PROBE = {"protocol": 1, "ok": True, "present": True, "version": "fixture-1",
         "dependency_fingerprint": "fixture-dependencies", "python_version": "fixture-python"}


def answer(engine="markitdown"):
    return {"protocol": 1, "ok": True, "engine": engine, "version": "fixture-1",
            "raw": {"text_content": "# Control\n\nDato: 42"}, "markdown": "# Control\n\nDato: 42",
            "units": [{"number": 1, "markdown": "# Control\n\nDato: 42"}], "partial": True}


class LocalEngineTests(TestCase):
    def setUp(self):
        self.temp = workspace_temp()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.root = self.folder / "biblioteca"
        self.source = self.folder / "control.html"
        self.source.write_text("<h1>Control</h1><p>Dato: 42</p>", encoding="utf-8")
        self.config = self.folder / "instalacion" / "motores.local.json"
        self.config.parent.mkdir()
        self.addCleanup(patch.stopall)
        patch.object(le, "CONFIG_PATH", self.config).start()
        patch.dict(os.environ, {}, clear=True).start()

    def convert_fake(self, value=None, engine="markitdown", **options):
        payload = answer(engine) if value is None else value
        with patch.object(le, "_probe", return_value=PROBE), patch.object(le, "_run_worker", return_value=payload):
            return le.convert_with_engine(self.root, self.source, engine, python=sys.executable, **options)

    def test_publish_preserves_raw_original_and_partial_warning(self):
        result = self.convert_fake()
        output = Path(result["output"])
        doc = load_json((output / "document.json").read_bytes())
        self.assertEqual((output / "fuente/original.html").read_bytes(), self.source.read_bytes())
        self.assertEqual(load_json((output / "crudos/markitdown.json").read_bytes()), answer()["raw"])
        self.assertEqual(doc["units"][0]["locator"], {"kind": "document", "value": self.source.name})
        self.assertEqual(result["source_sha256"], digest(self.source.read_bytes()))
        self.assertFalse(result["semantic_verified"])
        self.assertFalse(result["offline_verified"])
        self.assertTrue(verify_artifacts(output))
        self.assertTrue(any("parcial" in w for w in result["warnings"]))

    def test_cache_does_not_repeat_engine_and_dependency_change_invalidates(self):
        with patch.object(le, "_probe", return_value=PROBE) as probe, patch.object(le, "_run_worker", return_value=answer()) as run:
            first = le.convert_with_engine(self.root, self.source, "markitdown", python=sys.executable)
            second = le.convert_with_engine(self.root, self.source, "markitdown", python=sys.executable)
            self.assertEqual(run.call_count, 1)
            self.assertEqual(first["id"], second["id"])
            self.assertTrue(second["reused"])
            probe.return_value = {**PROBE, "dependency_fingerprint": "new-dependencies"}
            third = le.convert_with_engine(self.root, self.source, "markitdown", python=sys.executable)
            self.assertEqual(run.call_count, 2)
            self.assertNotEqual(first["id"], third["id"])

    def test_docling_raw_list_and_missing_page_not_invented(self):
        import pymupdf
        self.source = self.folder / "control.pdf"
        with pymupdf.open() as pdf:
            for _ in range(2):
                pdf.new_page().insert_text((72, 72), "Control")
            pdf.save(self.source)
        models = self.folder / "models"
        (models / "EasyOcr").mkdir(parents=True)
        (models / "EasyOcr" / "fixture.bin").write_bytes(b"synthetic-not-a-model")
        payload = answer("docling")
        payload["raw"] = [{"page_range": [1, 2], "document": {"pages": {"1": {}}}}]
        result = self.convert_fake(payload, "docling", pages=[1, 2], models_path=models)
        self.assertEqual(result["missing"], [2])
        self.assertEqual(result["status"], "parcial")
        output = Path(result["output"])
        self.assertIsInstance(load_json((output / "crudos/docling.json").read_bytes()), list)

    def test_reject_invalid_worker_units_and_raw(self):
        for mutate in (lambda x: x.update(raw="wrong"), lambda x: x.update(units=[{"number": True, "markdown": "x"}]),
                       lambda x: x["units"].append(x["units"][0]), lambda x: x.update(engine="marker")):
            with self.subTest(mutate=mutate):
                value = answer()
                mutate(value)
                with self.assertRaises(le.LocalEngineError) as failure:
                    self.convert_fake(value)
                self.assertEqual(failure.exception.code, "worker_protocol")

    def test_missing_runtime_and_assets_stop_before_worker(self):
        with patch.object(le, "_run_worker") as worker:
            with self.assertRaises(le.LocalEngineError) as failure:
                le.convert_with_engine(self.root, self.source, "markitdown", python=self.folder / "python.exe")
            self.assertEqual(failure.exception.code, "python_missing")
            with self.assertRaises(le.LocalEngineError) as failure:
                le.convert_with_engine(self.root, self.source, "docling", python=sys.executable)
            self.assertEqual(failure.exception.code, "assets_missing")
            worker.assert_not_called()

    def test_url_audio_and_artificial_pages_rejected(self):
        with self.assertRaises(le.LocalEngineError):
            le.convert_with_engine(self.root, "https://example.invalid/file.html", "markitdown", python=sys.executable)
        for filename, pages in (("audio.mp3", None), ("control.html", [1])):
            with self.subTest(filename=filename), self.assertRaises(ValueError):
                le._inventory(Path(filename), "markitdown", pages)

    def test_configuration_is_portable_and_preserves_other_engine(self):
        executable = self.config.parent / "workers" / "markitdown" / "python.exe"
        executable.parent.mkdir(parents=True)
        executable.write_bytes(b"fake-runtime-not-executed")
        le.configure_engine("markitdown", executable)
        le.configure_engine("docling", sys.executable)
        value = load_json(self.config.read_bytes())
        self.assertEqual(value["engines"]["markitdown"]["python"], "workers/markitdown/python.exe")
        self.assertIn("docling", value["engines"])
        moved = self.folder / "otra_pc"
        moved.mkdir()
        (moved / "motores.local.json").write_bytes(self.config.read_bytes())
        executable2 = moved / "workers/markitdown/python.exe"
        executable2.parent.mkdir(parents=True)
        executable2.write_bytes(executable.read_bytes())
        with patch.object(le, "CONFIG_PATH", moved / "motores.local.json"):
            self.assertEqual(le.engine_python("markitdown"), executable2.resolve())
            self.assertEqual(le.engine_configuration("markitdown")["python"], str(executable2.resolve()))

    def test_local_models_configuration_is_relative_and_network_paths_are_rejected(self):
        models = self.config.parent / "modelos/docling"
        models.mkdir(parents=True)
        le.configure_engine("docling", sys.executable, models)
        value = load_json(self.config.read_bytes())
        self.assertEqual(value["engines"]["docling"]["models_path"], "modelos/docling")
        self.assertEqual(le.engine_configuration("docling")["models_path"], str(models.resolve()))
        with self.assertRaises(ValueError):
            le.configure_engine("docling", sys.executable, "//example.invalid/share/models")

    def test_explicit_invalid_environment_does_not_fall_back_to_saved_python(self):
        le.configure_engine("markitdown", sys.executable)
        with patch.dict(os.environ, {"SISTEMA_MD_MARKITDOWN_PYTHON": str(self.folder / "missing/python.exe")}):
            with self.assertRaises(le.LocalEngineError):
                le.engine_python("markitdown")
            self.assertEqual(le.engine_python("markitdown", sys.executable), Path(sys.executable).resolve())

    def test_invalid_configuration_and_relative_escape_are_rejected(self):
        self.config.write_bytes(json_bytes({"version": 1, "engines": {"markitdown": {"api_key": "synthetic"}}}))
        with self.assertRaises(ValueError):
            le.engine_configuration("markitdown")
        self.config.write_bytes(json_bytes({"version": 1, "engines": {"markitdown": {"python": "../python.exe"}}}))
        with self.assertRaises(ValueError):
            le.engine_configuration("markitdown")

    def test_environment_override_does_not_silently_ignore_selected_runtime(self):
        python = self.config.parent / 'python.exe'
        python.write_bytes(b'fake-not-executed')
        with patch.dict(os.environ, {'SISTEMA_MD_MARKITDOWN_PYTHON':sys.executable}), \
             self.assertRaisesRegex(ValueError,'SISTEMA_MD_MARKITDOWN_PYTHON'):
            le.configure_engine('markitdown', python)
        self.assertFalse(self.config.exists())

    def test_environment_does_not_inherit_credentials_or_pythonpath(self):
        with patch.dict(os.environ, {"API_KEY": "synthetic", "OPENAI_API_KEY": "synthetic", "PYTHONPATH": "bad"}):
            env = le._environment(self.folder)
        self.assertNotIn("API_KEY", env)
        self.assertNotIn("OPENAI_API_KEY", env)
        self.assertNotIn("PYTHONPATH", env)
        self.assertEqual(env["HF_HUB_OFFLINE"], "1")

    def test_process_timeout_kills_child(self):
        process = Mock()
        process.stdin = io.BytesIO()
        process.poll.return_value = None
        with patch.object(le.subprocess, "Popen", return_value=process) as start, patch.object(le.time, "monotonic", side_effect=[0, 2]):
            with self.assertRaises(le.LocalEngineError) as failure:
                le._run_worker(sys.executable, {}, self.folder, 1)
        self.assertEqual(failure.exception.code, "worker_timeout")
        process.kill.assert_called_once()
        self.assertEqual(start.call_args.args[0][1:3], ["-I", "-B"])
        self.assertFalse(start.call_args.kwargs["shell"])

    def test_worker_output_limit_is_not_truncated_and_published(self):
        process = Mock()
        process.stdin = io.BytesIO()
        process.poll.return_value = None
        def start(*args, **kwargs):
            kwargs["stdout"].write(b"x" * 30)
            kwargs["stdout"].flush()
            return process
        with patch.object(le.subprocess, "Popen", side_effect=start), patch.object(le, "MAX_OUTPUT_BYTES", 10):
            with self.assertRaises(le.LocalEngineError) as failure:
                le._run_worker(sys.executable, {}, self.folder, 1)
        self.assertEqual(failure.exception.code, "worker_output_limit")
        process.kill.assert_called_once()

    def test_real_probe_does_not_import_engine(self):
        result = le._probe("markitdown", Path(sys.executable), self.folder)
        self.assertIsNone(result["importable"])
        self.assertIsNone(result["models_ready"])
        self.assertIn("dependency_fingerprint", result)

    def test_bad_worker_protocol_returns_safe_error(self):
        result = subprocess.run([sys.executable, "-I", "-B", str(le.WORKER)],
            input=b'{"protocol":1,"protocol":1}', capture_output=True, timeout=20, env=le._environment(self.folder))
        value = load_json(result.stdout)
        self.assertFalse(value["ok"])
        self.assertEqual(value["code"], "worker_failed")

    def test_audit_guard_blocks_network_and_subprocess_before_effect(self):
        program = """import importlib.util,socket,subprocess,sys
spec=importlib.util.spec_from_file_location('worker',sys.argv[1])
w=importlib.util.module_from_spec(spec); spec.loader.exec_module(w); w.install_local_guard()
for name,call in [('network_blocked',lambda: socket.getaddrinfo('example.invalid',443)), ('process_blocked',lambda: subprocess.Popen([sys.executable,'-V']))]:
    try: call()
    except w.GuardError as e: assert str(e)==name
    else: raise AssertionError(name)
print('blocked-before-effect')
"""
        result = subprocess.run([sys.executable, "-I", "-B", "-c", program, str(le.WORKER)],
            capture_output=True, timeout=20, env=le._environment(self.folder))
        self.assertEqual(result.returncode, 0, result.stderr.decode(errors="replace"))
        self.assertIn(b"blocked-before-effect", result.stdout)

    def test_archive_traversal_is_rejected(self):
        spec = importlib.util.spec_from_file_location("worker_test", le.WORKER)
        worker = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(worker)
        path = self.folder / "bad.docx"
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr("../escape", "bad")
        with self.assertRaisesRegex(ValueError, "archive_rejected"):
            worker._archive_safe(path)

    def test_real_markitdown_html_control_when_installed(self):
        python = le.WORKERS_ROOT / "markitdown/.venv" / ("Scripts/python.exe" if os.name == "nt" else "bin/python")
        if not python.is_file():
            self.skipTest("Worker MarkItDown no instalado: no se descarga en pruebas")
        self.source.write_text('<html><h1>Control local 42</h1><table><tr><th>Dato</th><th>Valor</th></tr>'
            '<tr><td>Cemento</td><td>123</td></tr></table><img src="https://example.invalid/x.png">'
            '<script>fetch("https://example.invalid/")</script></html>', encoding="utf-8")
        result = le.convert_with_engine(self.root, self.source, "markitdown", python=python, timeout=60)
        text = (Path(result["output"]) / "crudos/markitdown.md").read_text(encoding="utf-8")
        self.assertIn("Control local 42", text)
        self.assertIn("Cemento", text)
        self.assertIn("123", text)
        self.assertNotIn("fetch(", text)
        self.assertEqual(result["external_calls"], 0)
        self.assertFalse(result["offline_verified"])

    def test_real_markitdown_office_controls_when_installed(self):
        python = le.WORKERS_ROOT / "markitdown/.venv" / ("Scripts/python.exe" if os.name == "nt" else "bin/python")
        if not python.is_file():
            self.skipTest("Worker MarkItDown no instalado: no se descarga en pruebas")
        from docx import Document
        from openpyxl import Workbook
        from pptx import Presentation
        word = Document()
        word.add_heading("Control Word 42", 0)
        word.add_paragraph("Cemento 123")
        word.save(self.folder / "control.docx")
        book = Workbook()
        book.active.append(["Material", "Valor"])
        book.active.append(["Cemento", 123])
        book.save(self.folder / "control.xlsx")
        slides = Presentation()
        slide = slides.slides.add_slide(slides.slide_layouts[1])
        slide.shapes.title.text = "Control PowerPoint 42"
        slide.placeholders[1].text = "Cemento 123"
        slides.save(self.folder / "control.pptx")
        for suffix in (".docx", ".xlsx", ".pptx"):
            with self.subTest(suffix=suffix):
                source = self.folder / ("control" + suffix)
                result = le.convert_with_engine(self.root, source, "markitdown", python=python, timeout=60)
                text = (Path(result["output"]) / "crudos/markitdown.md").read_text(encoding="utf-8")
                self.assertIn("Cemento", text)
                self.assertIn("123", text)
                self.assertFalse(result["semantic_verified"])
                self.assertFalse(result["offline_verified"])

    def test_prepare_worker_default_only_prints_plan_and_refuses_overwrite(self):
        path = le.WORKER.with_name("prepare_markitdown.py")
        spec = importlib.util.spec_from_file_location("prepare_test", path)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        with patch.object(module.subprocess, "run") as run, patch("builtins.print"):
            self.assertEqual(module.main(["--python", sys.executable]), 0)
            run.assert_not_called()
            with patch.object(module, "TARGET", self.folder):
                self.assertEqual(module.main(["--python", sys.executable, "--install"]), 2)
            run.assert_not_called()

    def test_real_markitdown_epub_and_text_controls_when_installed(self):
        python = le.WORKERS_ROOT / "markitdown/.venv" / ("Scripts/python.exe" if os.name == "nt" else "bin/python")
        if not python.is_file():
            self.skipTest("Worker MarkItDown no instalado: no se descarga en pruebas")
        for suffix, text in ((".txt", "Control Cemento 123"), (".md", "# Control\n\nCemento 123"),
                             (".csv", "Material,Valor\nCemento,123\n")):
            (self.folder / ("control" + suffix)).write_text(text, encoding="utf-8")
        epub = self.folder / "control.epub"
        with zipfile.ZipFile(epub, "w") as archive:
            archive.writestr("mimetype", "application/epub+zip")
            archive.writestr("META-INF/container.xml", '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
            archive.writestr("OEBPS/content.opf", '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="id">fixture-control</dc:identifier><dc:title>Control</dc:title><dc:language>es</dc:language></metadata><manifest><item id="chap" href="chapter.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="chap"/></spine></package>')
            archive.writestr("OEBPS/chapter.xhtml", '<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Control</title></head><body><h1>Control</h1><p>Cemento 123</p></body></html>')
        for suffix in (".txt", ".md", ".csv", ".epub"):
            with self.subTest(suffix=suffix):
                result = le.convert_with_engine(self.root, self.folder / ("control" + suffix), "markitdown", python=python, timeout=60)
                text = (Path(result["output"]) / "crudos/markitdown.md").read_text(encoding="utf-8")
                self.assertIn("Cemento", text)
                self.assertIn("123", text)

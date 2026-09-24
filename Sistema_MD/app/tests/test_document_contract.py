"""Regresiones de F1: H01 (contrato), H02 (conector PPT), H03 (Excel grande).

Cada prueba reproduce el defecto tal como se midió el 2026-09-20 y exige el comportamiento
que pide PLAN_MAESTRO_IMPLEMENTACION.md (F1.02 a F1.06). Los controles sanos acompañan a
cada caso: un arreglo que rompe el flujo legítimo no es un arreglo.
"""

import importlib.util
import os
import zipfile
from pathlib import Path
from unittest import TestCase, mock, skipUnless

from conversion.documents import validate
from support import workspace_temp

HAS = {m: importlib.util.find_spec(m) is not None for m in ("openpyxl", "pptx")}


def _doc(**cambios) -> dict:
    base = {"schema_version": 1, "title": "Documento de prueba", "expected_units": [1],
            "units": [{"number": 1, "blocks": [
                {"id": "p1-b1", "kind": "paragraph", "text": "Contenido real.\n\n"}]}]}
    base.update(cambios)
    return base


class ContratoDocumento(TestCase):
    """H01: validate() exige el contrato, no confía en el productor."""

    def test_sin_titulo_se_rechaza(self):
        doc = _doc()
        del doc["title"]
        qa = validate(doc)
        self.assertFalse(qa["structural_ok"])
        self.assertTrue(any("ítulo" in e for e in qa["errors"]), qa["errors"])

    def test_titulo_vacio_se_rechaza(self):
        qa = validate(_doc(title="   "))
        self.assertFalse(qa["structural_ok"])

    def test_encabezado_fuera_de_rango_se_rechaza(self):
        doc = _doc(units=[{"number": 1, "blocks": [
            {"id": "p1-b1", "kind": "heading", "level": 42, "text": "# Título\n\n"}]}])
        qa = validate(doc)
        self.assertFalse(qa["structural_ok"])
        self.assertTrue(any("fuera de rango" in e for e in qa["errors"]), qa["errors"])

    def test_encabezado_valido_pasa(self):
        doc = _doc(units=[{"number": 1, "blocks": [
            {"id": "p1-b1", "kind": "heading", "level": 2, "text": "## Título\n\n"}]}])
        self.assertTrue(validate(doc)["structural_ok"])

    def test_encabezado_sin_nivel_sigue_siendo_valido(self):
        """Control sano: las respuestas de IA traen kind+text sin level. No deben romperse."""
        doc = _doc(units=[{"number": 1, "blocks": [
            {"id": "p1-b1", "kind": "heading", "text": "## Del motor\n\n"}]}])
        self.assertTrue(validate(doc)["structural_ok"])


@skipUnless(HAS["pptx"], "python-pptx no está instalado")
class ConectorPowerPoint(TestCase):
    """H02: un conector pegado a dos figuras no puede desaparecer en silencio."""

    def setUp(self):
        self.temp = workspace_temp()
        self.addCleanup(self.temp.cleanup)

    def _presentacion(self, pegar: bool) -> Path:
        from pptx import Presentation
        from pptx.util import Emu
        from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
        pr = Presentation()
        slide = pr.slides.add_slide(pr.slide_layouts[6])
        caja1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                       Emu(914400), Emu(914400), Emu(1828800), Emu(685800))
        caja1.text_frame.text = "Inicio"
        caja2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                       Emu(4572000), Emu(914400), Emu(1828800), Emu(685800))
        caja2.text_frame.text = "Aprobado"
        conector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                              Emu(2743200), Emu(1257300), Emu(4572000), Emu(1257300))
        if pegar:
            conector.begin_connect(caja1, 3)
            conector.end_connect(caja2, 1)
        destino = self.temp.path / ("flujo_%s.pptx" % ("pegado" if pegar else "suelto"))
        pr.save(destino)
        return destino

    def test_conector_pegado_queda_en_el_markdown(self):
        from conversion.native import pptx_texto
        ruta = self._presentacion(pegar=True)
        md, _visual, _visuales = pptx_texto(ruta, self.temp.path / "salida_pegado")
        self.assertIn("Conector", md, md)
        self.assertIn("Inicio", md)
        self.assertIn("Aprobado", md)
        self.assertNotIn("Inicio → Aprobado", md,
                         "Pegar extremos no declara el sentido de una flecha")
        self.assertIn("relación geométrica", md)

    def test_conector_suelto_marca_la_diapositiva(self):
        from conversion.native import pptx_texto
        ruta = self._presentacion(pegar=False)
        _md, _visual, visuales = pptx_texto(ruta, self.temp.path / "salida_suelto")
        self.assertIn(1, visuales, "un conector sin extremos debe quedar pendiente de revisión")


@skipUnless(HAS["openpyxl"], "openpyxl no está instalado")
class ExcelGrandeCeldaVacia(TestCase):
    """H03: en la rama de solo lectura los huecos llegan como EmptyCell, sin ``column``."""

    def setUp(self):
        self.temp = workspace_temp()
        self.addCleanup(self.temp.cleanup)

    def _libro(self) -> Path:
        from openpyxl import Workbook
        wb = Workbook()
        hoja = wb.active
        hoja.title = "Datos"
        hoja["A1"] = "Partida"
        hoja["C1"] = "Total"          # B1 vacía: el hueco intermedio del defecto
        hoja["A2"] = "Excavación"
        hoja["C2"] = 120
        hoja["C3"] = "=C2*2"
        destino = self.temp.path / "grande.xlsx"
        wb.save(destino)
        return destino

    def test_celda_intermedia_vacia_en_hoja_grande(self):
        from conversion import native
        ruta = self._libro()
        real = os.path.getsize

        def grande(path, *a, **k):          # fuerza la rama rich=False sin crear 30 MB
            return 40_000_000 if str(path) == str(ruta) else real(path, *a, **k)

        with mock.patch.object(native.os.path, "getsize", grande):
            resultado = native.excel_tres_capas(ruta, self.temp.path / "salida")
        self.assertIsNotNone(resultado)

    def test_control_sano_archivo_normal(self):
        from conversion import native
        resultado = native.excel_tres_capas(self._libro(), self.temp.path / "salida_normal")
        self.assertIsNotNone(resultado)

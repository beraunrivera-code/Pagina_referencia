import json
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from support import workspace_temp  # noqa: E402
from conversion.pipeline import convert_native  # noqa: E402
from conversion.storage import verify_artifacts  # noqa: E402
from conversion.visor import build_view, write_view  # noqa: E402


class ExcelStructureTests(unittest.TestCase):

    def setUp(self):
        self.temp = workspace_temp()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)

    def _workbook(self):
        import openpyxl
        from openpyxl.worksheet.table import Table

        workbook = openpyxl.Workbook()
        aux = workbook.active
        aux.title = "AUX"
        for row in (("Código", "Nombre", None, "Sigla", "Significado"),
                    ("A", "Arquitectura", None, "OT", "Oficina técnica"),
                    ("B", "BIM", None, "LB", "Línea base")):
            aux.append(row)
        aux.add_table(Table(displayName="Catalogo", ref="A1:B3"))
        aux.add_table(Table(displayName="Leyenda", ref="D1:E3"))

        main = workbook.create_sheet("Principal")
        main.merge_cells("E1:H1"); main["E1"] = "USUARIO A OFICINA TÉCNICA"
        main.merge_cells("I1:L1"); main["I1"] = "OFICINA TÉCNICA A PROCURA"
        headers = ["CÓDIGO", "NOMBRE", "ESTADO", "ALCANCE", "PROYECCIÓN A", "REAL A",
                   "RETRASO A", "S1", "PROYECCIÓN B", "REAL B", "RETRASO B", "S2", "NOTA"]
        for column, value in enumerate(headers, 1):
            main.cell(3, column, value)
        for row, code in ((4, "P-01"), (5, "P-02")):
            for column in range(1, 14):
                main.cell(row, column, code if column == 1 else f"D{row}-{column}")
        main.add_table(Table(displayName="Seguimiento", ref="A3:M5"))
        hidden = workbook.create_sheet("Cálculos")
        hidden.sheet_state = "hidden"
        hidden["A1"], hidden["B1"] = "Factor", 2
        source = self.base / "libro.xlsx"
        workbook.save(source)
        return source

    def test_excel_separa_tablas_y_grupos_sin_perder_rejilla(self):
        result = convert_native(self.base / "out", self._workbook())
        folder = Path(result["output"])
        markdown = (folder / "documento.md").read_text(encoding="utf-8")
        self.assertIn("## Hojas principales", markdown)
        self.assertIn("### Catalogo", markdown)
        self.assertIn("### Leyenda", markdown)
        self.assertIn("### Seguimiento", markdown)
        self.assertIn("USUARIO A OFICINA TÉCNICA", markdown)
        self.assertIn("Tabla original de 13 columnas", markdown)
        self.assertIn("## Hojas auxiliares", markdown)
        self.assertTrue((folder / "derivados" / "rejilla_original.md").is_file())
        self.assertTrue((folder / "derivados" / "celdas.csv").is_file())
        self.assertTrue((folder / "derivados" / "estructura.json").is_file())
        self.assertTrue(verify_artifacts(folder))

        structure = json.loads((folder / "derivados" / "estructura.json").read_text(encoding="utf-8"))
        main = next(section for group in structure["groups"] for section in group["sections"]
                    if section["title"] == "Principal")
        table = next(block for block in main["blocks"] if block["kind"] == "table_group")
        covered = {column for part in table["parts"] for column in part["columns"]
                   if column not in part["repeated_columns"]}
        self.assertEqual(covered, set(range(1, 14)))

    def test_visor_usa_titulo_real_y_ofrece_descargas(self):
        result = convert_native(self.base / "out", self._workbook())
        folder = Path(result["output"])
        html = build_view(folder)
        self.assertIn("libro · Sistema MD", html)
        self.assertNotIn("metadatos no disponibles", html)
        self.assertIn("Mapa de lectura", html)
        self.assertIn("Rejilla original", html)
        self.assertIn("download", html)
        view = write_view(folder)
        self.assertTrue(view.is_file())
        self.assertTrue(verify_artifacts(folder))

    # ---------------------------------------------------------------- ADN de 3 capas
    def _workbook_adn(self):
        import io
        import zipfile
        import openpyxl
        from openpyxl.drawing.image import Image
        from openpyxl.worksheet.table import Table
        from PIL import Image as Pil

        workbook = self._workbook()
        book = openpyxl.load_workbook(workbook)
        hidden = book["Cálculos"]
        hidden["A2"], hidden["B2"] = "Total", "=B1*3"
        png = io.BytesIO()
        Pil.new("RGB", (4, 4), "red").save(png, format="PNG")
        png.seek(0)
        book["AUX"].add_image(Image(png), "G2")
        precios = book.create_sheet("Precios")
        for row in (("Código", "Descripción", "Total"), ("P-1", "Cemento; saco", 12.5)):
            precios.append(row)
        precios.add_table(Table(displayName="Precios", ref="A1:C2"))
        book.save(workbook)
        # Un medio que openpyxl no genera pero Excel sí guarda: debe aislarse igual.
        with zipfile.ZipFile(workbook, "a") as archive:
            archive.writestr("xl/media/image9.wmf", b"WMF-fixture")
        return workbook

    def test_capa1_encabezados_conservan_letra_de_columna(self):
        result = convert_native(self.base / "out", self._workbook_adn())
        folder = Path(result["output"])
        markdown = (folder / "documento.md").read_text(encoding="utf-8")
        self.assertIn("| Código [A] | Descripción [B] | Total [C] |", markdown)
        self.assertIn("| Sigla [D] | Significado [E] |", markdown)
        self.assertIn("CÓDIGO [A]", markdown)          # partes de una tabla ancha dividida
        self.assertIn("PROYECCIÓN B [I]", markdown)
        structure = json.loads((folder / "derivados" / "estructura.json").read_text(encoding="utf-8"))
        precios = next(section for group in structure["groups"] for section in group["sections"]
                       if section["title"] == "Precios")
        self.assertEqual(precios["blocks"][0]["column_letters"], ["A", "B", "C"])
        self.assertEqual(precios["blocks"][0]["headers"], ["Código", "Descripción", "Total"])

    def test_capa1_hoja_oculta_siempre_marcada(self):
        result = convert_native(self.base / "out", self._workbook_adn())
        folder = Path(result["output"])
        for name in ("documento.md", "derivados/rejilla_original.md"):
            text = (folder / name).read_text(encoding="utf-8")
            self.assertIn("## Hoja: Cálculos *(oculta)*\n", text, name)
            self.assertNotIn("## Hoja: Precios *(oculta)*", text, name)

    def test_capa2_formulas_csv_utf8_sin_bom_y_punto_y_coma(self):
        import csv
        result = convert_native(self.base / "out", self._workbook_adn())
        raw = (Path(result["output"]) / "derivados" / "formulas.csv").read_bytes()
        self.assertFalse(raw.startswith(b"\xef\xbb\xbf"))
        text = raw.decode("utf-8")
        self.assertTrue(text.startswith("hoja;celda;formula;valor_cacheado\r\n"))
        rows = list(csv.reader(text.splitlines(), delimiter=";"))
        self.assertIn(["Cálculos", "B2", "=B1*3", ""], rows)

    def test_capa3_medios_aislados_en_imagenes_y_pendiente(self):
        result = convert_native(self.base / "out", self._workbook_adn())
        folder = Path(result["output"])
        self.assertTrue((folder / "imagenes" / "image1.png").is_file())
        self.assertEqual((folder / "imagenes" / "image9.wmf").read_bytes(), b"WMF-fixture")
        self.assertFalse(any(path.suffix in {".png", ".wmf"} for path in (folder / "derivados").iterdir()))
        markdown = (folder / "documento.md").read_text(encoding="utf-8")
        self.assertIn("[PENDIENTE: 2 imágenes embebidas guardadas en imagenes/ del paquete", markdown)
        self.assertTrue(verify_artifacts(folder))


if __name__ == "__main__":
    unittest.main()

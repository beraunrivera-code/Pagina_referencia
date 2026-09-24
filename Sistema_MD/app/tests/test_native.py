"""Adaptadores locales de Word, Excel, PowerPoint, Visio y CAD: 0 llamadas a IA.

Los archivos de prueba se generan aquí mismo con las librerías del adaptador; DWG binario no se
sintetiza (requiere ODA sobre un plano real) y se cubre en la prueba de extremo a extremo del CLI.
"""

import base64
import importlib.util
import zipfile
from pathlib import Path
from unittest import TestCase, skipUnless

from conversion.documents import load_json
from conversion.pipeline import NATIVE_KINDS, convert_native, detect
from conversion.storage import verify_artifacts
from conversion.workflow import WorkQueue, action_for
from support import workspace_temp

HAS = {m: importlib.util.find_spec(m) is not None for m in ("docx", "openpyxl", "pptx", "ezdxf")}
PNG_1x1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")


def markdown_of(result: dict) -> str:
    return (Path(result["output"]) / "documento.md").read_text(encoding="utf-8")


def manifest_of(result: dict) -> dict:
    return load_json((Path(result["output"]) / "manifest.json").read_bytes())


class NativeAdapterTests(TestCase):
    def setUp(self):
        temporary = workspace_temp()
        self.addCleanup(temporary.cleanup)
        self.base = Path(temporary.name)
        self.output = self.base / "memoria"

    def test_native_kinds_are_ready_and_legacy_office_stays_pending(self):
        for kind in NATIVE_KINDS:
            self.assertEqual(action_for(kind), ("listo", "Convertir local"))
        self.assertEqual(action_for("office-legacy")[0], "pendiente")
        legacy = self.base / "viejo.xls"
        legacy.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 64)
        queue = WorkQueue(self.output)
        queue.add(legacy)
        self.assertEqual(queue.items[0]["status"], "pendiente")
        with self.assertRaisesRegex(ValueError, "todavía no está implementado"):
            queue.process(str(legacy.resolve()))
        with self.assertRaisesRegex(ValueError, "todavía no está implementado"):
            convert_native(self.output, legacy)

    @skipUnless(HAS["docx"], "python-docx no instalado")
    def test_docx_text_tables_and_images_become_a_verified_package(self):
        import docx
        document = docx.Document()
        document.add_heading("Alcance", 1)
        document.add_paragraph("Dato verificable.")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text, table.rows[0].cells[1].text = "A", "B"
        image = self.base / "pixel.png"
        image.write_bytes(PNG_1x1)
        document.add_picture(str(image))
        source = self.base / "nota.docx"
        document.save(source)
        queue = WorkQueue(self.output)
        queue.add(source)
        self.assertEqual(queue.items[0]["status"], "listo")
        result = queue.process(str(source.resolve()))
        self.assertEqual(result["status"], "revisar")
        self.assertTrue(result["structural_ok"])
        self.assertFalse(result["reused"])
        text = markdown_of(result)
        for expected in ("## Alcance", "Dato verificable.", "| A | B |", "## Testigo",
                         "[PENDIENTE: 1 imágenes embebidas"):
            self.assertIn(expected, text)
        self.assertIn("imagenes/img01.png", manifest_of(result)["files"])
        self.assertTrue(any("declaró dudas" in w for w in result["warnings"]))
        self.assertTrue(verify_artifacts(Path(result["output"])))
        self.assertEqual(queue.summary()["completed"], 1)
        self.assertTrue(queue.process(str(source.resolve()))["reused"])

    @skipUnless(HAS["openpyxl"], "openpyxl no instalado")
    def test_xlsx_keeps_formulas_as_asset_even_when_extension_lies(self):
        import openpyxl
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Datos"
        sheet["A1"], sheet["B1"], sheet["A2"], sheet["B2"] = "Partida", 2, "Total", "=B1*3"
        honest = self.base / "libro.xlsx"
        liar = self.base / "mentira.pptx"  # medido en el lote real: un Excel llamado .pptx
        workbook.save(honest)
        workbook.save(liar)
        self.assertEqual(detect(liar), "xlsx")
        result = convert_native(self.output, liar)
        self.assertEqual(result["status"], "revisar")
        text = markdown_of(result)
        self.assertIn("# mentira.pptx\n", text)  # el nombre real, no el de la copia temporal
        self.assertNotIn("archivo.xlsx", text)
        self.assertIn("## Hoja: Datos", text)
        self.assertIn("| Partida | 2 |", text)
        self.assertIn("derivados/formulas.csv", manifest_of(result)["files"])
        formulas = (Path(result["output"]) / "derivados" / "formulas.csv").read_text(encoding="utf-8")
        self.assertIn("=B1*3", formulas)
        document = load_json((Path(result["output"]) / "document.json").read_bytes())
        self.assertEqual(document["input_kind"], "xlsx")
        self.assertEqual(document["title"], "mentira")
        self.assertTrue(liar.exists())  # el nativo no se toca ni se renombra

    @skipUnless(HAS["pptx"], "python-pptx no instalado")
    def test_pptx_text_and_notes_declare_visual_slides(self):
        from pptx import Presentation
        from pptx.util import Inches
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Flujo de cambios"
        slide.placeholders[1].text = "Paso uno"
        slide.notes_slide.notes_text_frame.text = "Recordar el RFI"
        image = self.base / "pixel.png"
        image.write_bytes(PNG_1x1)
        visual = presentation.slides.add_slide(presentation.slide_layouts[6])
        visual.shapes.add_picture(str(image), Inches(1), Inches(1))
        source = self.base / "deck.pptx"
        presentation.save(source)
        result = convert_native(self.output, source)
        text = markdown_of(result)
        for expected in ("## Diapositiva 1", "Flujo de cambios", "Paso uno",
                         "> **Notas del orador:** Recordar el RFI", "## Diapositiva 2",
                         "[PENDIENTE: 1 diapositivas con contenido visual (2)", "## Testigo"):
            self.assertIn(expected, text)
        self.assertTrue(any("declaró dudas" in w for w in result["warnings"]))

    def _write_vsdx(self, path: Path, with_connectors: bool) -> None:
        page = ('<?xml version="1.0" encoding="utf-8"?>'
                '<PageContents xmlns="http://schemas.microsoft.com/office/visio/2012/main"><Shapes>'
                '<Shape ID="1" Type="Shape"><Text>Inicio</Text></Shape>'
                '<Shape ID="2" Type="Shape"><Text>Fin &gt; 15 dc</Text></Shape>'
                '<Shape ID="3" Type="Shape"><Text>Sí</Text></Shape></Shapes>')
        if with_connectors:
            page += ('<Connects><Connect FromSheet="3" FromCell="BeginX" ToSheet="1"/>'
                     '<Connect FromSheet="3" FromCell="EndX" ToSheet="2"/></Connects>')
        page += "</PageContents>"
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr("visio/document.xml", "<VisioDocument/>")
            archive.writestr("visio/pages/page1.xml", page)

    def test_vsdx_rebuilds_flow_from_connectors_or_declares_none(self):
        flow = self.base / "flujo.vsdx"
        self._write_vsdx(flow, with_connectors=True)
        self.assertEqual(detect(flow), "vsdx")
        result = convert_native(self.output, flow)
        text = markdown_of(result)
        self.assertIn("> - Inicio —«Sí»— → Fin > 15 dc", text)
        self.assertIn("## Testigo\n> «Inicio»", text)
        self.assertNotIn("sin conectores legibles: el diagrama NO se reconstruyó", result["warnings"])
        boxes = self.base / "cajas.vsdx"
        self._write_vsdx(boxes, with_connectors=False)
        result = convert_native(self.output, boxes)
        self.assertEqual(result["status"], "revisar")
        self.assertIn("Sin conectores legibles", markdown_of(result))
        self.assertIn("sin conectores legibles: el diagrama NO se reconstruyó", result["warnings"])

    def test_dxf_is_detected_by_content_not_extension(self):
        ascii_dxf = self.base / "plano.txt"
        ascii_dxf.write_bytes(b"999\r\ncomentario\r\n  0\r\nSECTION\r\n  2\r\nHEADER\r\n  0\r\nENDSEC\r\n  0\r\nEOF\r\n")
        self.assertEqual(detect(ascii_dxf), "dxf")
        binary_dxf = self.base / "plano.bin"
        binary_dxf.write_bytes(b"AutoCAD Binary DXF\r\n\x1a\x00" + b"\x00" * 32)
        self.assertEqual(detect(binary_dxf), "dxf")
        prose = self.base / "nota.txt"
        prose.write_text("0 SECTION no es un DXF si no va por líneas", encoding="utf-8")
        self.assertEqual(detect(prose), "txt")

    @skipUnless(HAS["ezdxf"], "ezdxf no instalado")
    def test_dxf_is_measured_not_transcribed(self):
        import ezdxf
        drawing = ezdxf.new()
        space = drawing.modelspace()
        space.add_line((0, 0, 0), (3, 4, 0), dxfattribs={"layer": "E-TIERRA"})
        space.add_text("TABLERO TG-01 ALIMENTADOR PRINCIPAL", dxfattribs={"layer": "TEXTO"})
        source = self.base / "plano.dxf"
        drawing.saveas(source)
        self.assertEqual(detect(source), "dxf")
        result = convert_native(self.output, source)
        self.assertEqual(result["status"], "revisar")
        text = markdown_of(result)
        for expected in ("| Entidades | 2 |", "| E-TIERRA | 5.00 |", "- TABLERO TG-01 ALIMENTADOR PRINCIPAL",
                         "## Testigo\n> «TABLERO TG-01 ALIMENTADOR PRINCIPAL»"):
            self.assertIn(expected, text)
        files = manifest_of(result)["files"]
        self.assertIn("derivados/dwg_geometria.csv", files)
        self.assertIn("derivados/dwg_textos.csv", files)
        self.assertTrue(verify_artifacts(Path(result["output"])))

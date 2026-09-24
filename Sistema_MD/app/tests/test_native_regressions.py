"""Casos mínimos de fidelidad local y límites; no certifican otros documentos."""

import csv
import importlib.util
import sys
import zipfile
from types import SimpleNamespace
from unittest import TestCase, mock, skipUnless
from xml.etree import ElementTree as ET

from support import workspace_temp


@skipUnless(importlib.util.find_spec("pptx"), "python-pptx no instalado")
class ConnectorDirection(TestCase):
    def setUp(self):
        self.temp = workspace_temp()
        self.addCleanup(self.temp.cleanup)

    def extract(self, head=None, tail=None):
        from pptx import Presentation
        from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
        from pptx.oxml.xmlchemy import OxmlElement
        from pptx.util import Inches
        from conversion.native import pptx_texto
        pr = Presentation()
        slide = pr.slides.add_slide(pr.slide_layouts[6])
        first = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
        first.text = "Origen geométrico"
        last = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5), Inches(1), Inches(2), Inches(1))
        last.text = "Extremo geométrico"
        connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(3), Inches(1.5), Inches(5), Inches(1.5))
        connector.begin_connect(first, 3)
        connector.end_connect(last, 1)
        line = connector._element.spPr.get_or_add_ln()
        for tag, value in (("headEnd", head), ("tailEnd", tail)):
            if value is not None:
                node = OxmlElement("a:" + tag)
                node.set("type", value)
                line.append(node)
        path = self.temp.path / "conector.pptx"
        pr.save(path)
        return pptx_texto(path, self.temp.path / "salida")

    def test_attached_endpoints_do_not_invent_a_flow_arrow(self):
        md, _, pending = self.extract()
        self.assertNotIn("Origen geométrico → Extremo geométrico", md)
        self.assertIn("relación geométrica", md)
        self.assertIn(1, pending)

    def test_explicit_no_arrow_stays_geometric(self):
        md, _, _ = self.extract("none", "none")
        self.assertNotIn("→", md)
        self.assertIn("sin flechas explícitas", md)

    def test_begin_arrow_reverses_the_reading(self):
        md, _, _ = self.extract("triangle", "none")
        self.assertIn("Extremo geométrico → Origen geométrico", md)

    def test_end_arrow_keeps_the_reading(self):
        md, _, _ = self.extract("none", "triangle")
        self.assertIn("Origen geométrico → Extremo geométrico", md)

    def test_both_arrows_do_not_become_one_way(self):
        md, _, _ = self.extract("arrow", "triangle")
        self.assertIn("Origen geométrico ↔ Extremo geométrico", md)

    def test_diamond_is_not_a_directional_arrow(self):
        md, _, pending = self.extract("diamond", "none")
        self.assertNotIn("→", md)
        self.assertIn("headEnd=diamond", md)
        self.assertIn(1, pending)

    def test_new_native_version_does_not_reuse_previous_conversion(self):
        from conversion import pipeline
        self.extract("none", "triangle")
        source = self.temp.path / "conector.pptx"
        data = self.temp.path / "datos"
        self.assertGreater(pipeline.NATIVE_VERSION, 4)
        with mock.patch.object(pipeline, "NATIVE_VERSION", 4):
            previous = pipeline.convert_native(data, source)
        current = pipeline.convert_native(data, source)
        self.assertFalse(current["reused"])
        self.assertNotEqual(previous["id"], current["id"])
        with mock.patch.object(pipeline, "pptx_texto", side_effect=AssertionError("no reprocesar")):
            self.assertTrue(pipeline.convert_native(data, source)["reused"])


@skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl no instalado")
class SparseExcelCoordinates(TestCase):
    def setUp(self):
        self.temp = workspace_temp()
        self.addCleanup(self.temp.cleanup)

    def _book(self, dimension=True):
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.title = "Datos"
        sheet["C3"] = "Inicio con huecos"
        sheet["F5"] = 0
        sheet["H5"] = False
        sheet["C7"] = "=1+1"
        sheet["H9"] = "=2+2"
        sheet["C11"] = "=2-2"
        sheet["H13"] = "=1=1"
        source = self.temp.path / "original.xlsx"
        book.save(source)
        book.close()
        # Modificación del fixture, nunca del documento del usuario: caché conocida 2.
        with zipfile.ZipFile(source) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        namespace = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        root = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
        for cell in root.findall(".//m:c", namespace):
            if cell.get("r") == "C7":
                cell.find("m:v", namespace).text = "2"
            elif cell.get("r") == "C11":
                cell.find("m:v", namespace).text = "0"
            elif cell.get("r") == "H13":
                cell.set("t", "b")
                cell.find("m:v", namespace).text = "1"
        if not dimension:
            root.remove(root.find("m:dimension", namespace))
        parts["xl/worksheets/sheet1.xml"] = ET.tostring(root)
        with zipfile.ZipFile(source, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, data in parts.items():
                archive.writestr(name, data)
        return source

    def test_streaming_matches_rich_with_holes_zero_boolean_and_formula_cache(self):
        from conversion import native
        for dimension in (True, False):
            with self.subTest(dimension=dimension):
                source = self._book(dimension)
                before = source.read_bytes()
                outputs = []
                for rich in (True, False):
                    dest = self.temp.path / ("rich" if rich else "stream")
                    with mock.patch.object(native.os.path, "getsize", return_value=1 if rich else 40_000_000):
                        md, count, _, structure = native.excel_tres_capas(source, dest)
                    self.assertEqual(count, 4)
                    self.assertIn("[SIN CACHÉ] =2+2", md)
                    section = structure["groups"][0]["sections"][0]
                    self.assertEqual(section["stats"]["nonempty_cells"], 7)
                    self.assertEqual(section["stats"]["rows_with_data"], 6)
                    refs = {item["source_ref"] for block in section["blocks"]
                            if block["kind"] == "notes" for item in block["items"]}
                    self.assertEqual(refs, {"C3:C3", "F5:F5", "H5:H5", "C7:C7", "H9:H9", "C11:C11", "H13:H13"})
                    with (dest / "celdas.csv").open(encoding="utf-8", newline="") as stream:
                        cells = {row["celda"]: row for row in csv.DictReader(stream, delimiter=";")}
                    with (dest / "formulas.csv").open(encoding="utf-8", newline="") as stream:
                        formulas = list(csv.DictReader(stream, delimiter=";"))
                    self.assertEqual(set(cells), {"C3", "F5", "H5", "C7", "H9", "C11", "H13"})
                    self.assertEqual(cells["F5"]["valor_original"], "0")
                    self.assertEqual(cells["H5"]["valor_original"], "False")
                    self.assertEqual(cells["C7"]["valor_original"], "=1+1")
                    self.assertEqual(cells["C7"]["valor_cacheado"], "2")
                    self.assertEqual(cells["H9"]["valor_cacheado"], "")
                    self.assertEqual(cells["C11"]["valor_cacheado"], "0")
                    self.assertEqual(cells["H13"]["valor_cacheado"], "True")
                    self.assertEqual(formulas, [
                        {"hoja": "Datos", "celda": "C7", "formula": "=1+1", "valor_cacheado": "2"},
                        {"hoja": "Datos", "celda": "H9", "formula": "=2+2", "valor_cacheado": ""},
                        {"hoja": "Datos", "celda": "C11", "formula": "=2-2", "valor_cacheado": "0"},
                        {"hoja": "Datos", "celda": "H13", "formula": "=1=1", "valor_cacheado": "True"},
                    ])
                    outputs.append(cells)
                self.assertEqual(outputs[0], outputs[1])
                self.assertEqual(before, source.read_bytes(), "el original debe quedar intacto")


class PdfPreparationResources(TestCase):
    def setUp(self):
        self.temp = workspace_temp()
        self.addCleanup(self.temp.cleanup)

    def test_accumulated_assets_are_limited_before_publication(self):
        from conversion import pipeline
        source = self.temp.path / "sintetico.pdf"
        source.write_bytes(b"%PDF-fixture-mocked")
        page = mock.Mock()
        page.rect = SimpleNamespace(width=72, height=72)
        page.get_pixmap.return_value.tobytes.return_value = b"\x89PNG\r\n\x1a\n" + b"x" * 692
        page.get_text.return_value = "Texto testigo"
        page.get_drawings.return_value = []
        page.get_images.return_value = []
        page.annots.return_value = []
        pdf = mock.MagicMock()
        pdf.needs_pass = False
        pdf.__len__.return_value = 2
        pdf.__getitem__.return_value = page
        pdf.__enter__.return_value = pdf
        fitz = SimpleNamespace(open=mock.Mock(return_value=pdf))
        with mock.patch.dict(sys.modules, {"fitz": fitz}), \
                mock.patch.object(pipeline, "MAX_PREPARED_BYTES", 1024, create=True), \
                mock.patch.object(pipeline, "prepared_cache", return_value=None), \
                mock.patch.object(pipeline, "publish") as publish:
            with self.assertRaisesRegex(ValueError, "memoria.*lotes"):
                pipeline.prepare_pdf(self.temp.path / "datos", source, [1, 2])
            publish.assert_not_called()

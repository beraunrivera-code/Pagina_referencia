import csv
import json
import threading
from pathlib import Path
from unittest import TestCase
from unittest.mock import patch
from contextlib import closing

from support import workspace_temp
from conversion.batches import plan_queue, run_batch, batch_status, database
from conversion.local_io import exclusive, check_archive
from conversion.workflow import WorkQueue
from conversion.pipeline import convert_native
from conversion.storage import verify_artifacts
from conversion.credentials import get_key, key_location


class CouplingTests(TestCase):
    def setUp(self):
        fixture = workspace_temp()
        self.addCleanup(fixture.cleanup)
        self.base = Path(fixture.name)
        self.root = self.base / 'datos'
        self.sources = self.base / 'fuentes'
        self.sources.mkdir()

    def source(self, name, body):
        path = self.sources / name
        path.write_text(body, encoding='utf-8')
        return path

    def plan(self):
        WorkQueue(self.root).add(self.sources)
        return plan_queue(self.root)

    def test_batch_deduplicates_and_resumes_without_reopening_completed_sources(self):
        self.source('a.txt', 'La misma evidencia')
        self.source('b.txt', 'La misma evidencia')
        self.source('c.txt', 'Otra evidencia')
        self.source('d.pdf', '%PDF-1.4\nno invocar IA')
        plan = self.plan()
        self.assertEqual(plan['routes'], {'local': 2, 'duplicado': 1, 'pendiente': 1})
        with patch('conversion.batches.convert_text', wraps=__import__('conversion.pipeline', fromlist=['convert_text']).convert_text) as convert:
            first = run_batch(self.root, plan['id'], limit=1)
            self.assertEqual(first['status'], 'pausado')
            done = run_batch(self.root, plan['id'])
            self.assertEqual(convert.call_count, 2)
        self.assertEqual(done['counts'], {'terminado': 3, 'pendiente': 1})
        self.assertEqual(done['external_calls'], 0)
        self.assertEqual(done['items'][0]['result']['id'], done['items'][1]['result']['id'])
        with patch('conversion.batches.file_hash', side_effect=AssertionError('No releer completados')):
            self.assertEqual(run_batch(self.root)['processed'], 0)

    def test_pause_and_exclusive_worker(self):
        self.source('a.txt', 'Dato')
        self.plan()
        stop = threading.Event(); stop.set()
        self.assertEqual(run_batch(self.root, stop=stop)['processed'], 0)
        with exclusive(self.root):
            with self.assertRaisesRegex(ValueError, 'Otro proceso'):
                run_batch(self.root)
        self.assertEqual(run_batch(self.root)['processed'], 1)

    def test_source_changed_is_mapped_and_does_not_loop(self):
        source = self.source('a.txt', 'Antes')
        self.plan()
        source.write_text('Después', encoding='utf-8')
        failed = run_batch(self.root)
        self.assertEqual(failed['counts'], {'error': 1})
        self.assertIn('cambió', failed['items'][0]['result']['error'])
        self.assertEqual(run_batch(self.root)['processed'], 0)

    def test_plan_is_idempotent_and_failed_zip_is_not_a_document(self):
        self.source('a.txt', 'Texto')
        (self.sources / 'bad.zip').write_bytes(b'PKbad')
        plan = self.plan()
        self.assertEqual(plan['id'], plan_queue(self.root)['id'])
        self.assertEqual(plan['items'][1]['route'], 'pendiente')

    def test_interrupted_local_item_can_resume_and_corruption_is_not_reused(self):
        self.source('a.txt', 'Texto')
        plan = self.plan()
        with closing(database(self.root)) as db, db:
            db.execute("UPDATE local_items SET state='en_curso' WHERE batch=?", (plan['id'],))
        done = run_batch(self.root)
        old = Path(done['items'][0]['result']['output'])
        (old / 'documento.md').write_text('alterado', encoding='utf-8')
        result = run_batch(self.root)
        new = Path(result['items'][0]['result']['output'])
        self.assertNotEqual(old, new)
        self.assertTrue(old.exists())
        self.assertTrue(verify_artifacts(new))

    def test_excel_preserves_row_450_formula_quotes_and_literal_pipe(self):
        import openpyxl
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet['A1'] = 'A|B'
        sheet['A450'] = '=IF(1=1,"sí","no")'
        source = self.sources / 'datos.pptx'
        workbook.save(source)
        result = convert_native(self.root, source)
        folder = Path(result['output'])
        with (folder/'derivados/formulas.csv').open(encoding='utf-8', newline='') as stream:
            formulas = list(csv.DictReader(stream, delimiter=';'))
        self.assertEqual(formulas[0]['celda'], 'A450')
        self.assertEqual(formulas[0]['formula'], '=IF(1=1,"sí","no")')
        md = (folder/'documento.md').read_text(encoding='utf-8')
        self.assertIn('A&#124;B', md)
        self.assertIn('[SIN CACHÉ]', md)
        with (folder/'derivados/celdas.csv').open(encoding='utf-8', newline='') as stream:
            cells = list(csv.DictReader(stream, delimiter=';'))
        self.assertEqual(cells[0]['valor_original'], 'A|B')
        renamed = self.sources / 'otro.xlsx'
        renamed.write_bytes(source.read_bytes())
        with patch('conversion.pipeline.excel_tres_capas', side_effect=AssertionError('No extraer otra vez')):
            reused = convert_native(self.root, renamed)
        self.assertTrue(reused['reused'])

    def test_word_preserves_interleaved_table(self):
        import docx
        document = docx.Document()
        document.add_paragraph('ANTES')
        document.add_table(rows=1, cols=1).cell(0,0).text = 'TABLA|DATO'
        document.add_paragraph('DESPUÉS')
        source = self.sources / 'word.bin'
        document.save(source)
        result = convert_native(self.root, source)
        md = (Path(result['output'])/'documento.md').read_text(encoding='utf-8')
        self.assertLess(md.index('ANTES'), md.index('TABLA&#124;DATO'))
        self.assertLess(md.index('TABLA&#124;DATO'), md.index('DESPUÉS'))
        self.assertFalse(any((self.root/'temporales').iterdir()))

    def test_duplicate_zip_members_are_rejected(self):
        import zipfile
        import warnings
        source = self.sources/'bad.docx'
        with warnings.catch_warnings():
            warnings.simplefilter('ignore')
            with zipfile.ZipFile(source, 'w') as archive:
                archive.writestr('word/document.xml', '<x/>')
                archive.writestr('word/document.xml', '<x/>')
        with self.assertRaisesRegex(ValueError, 'duplicados'):
            check_archive(source)

    def test_pptx_keeps_table_and_group_text(self):
        from pptx import Presentation
        from pptx.util import Inches
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        table = slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(2), Inches(1)).table
        table.cell(0, 0).text = 'TABLA|PPT'
        group = slide.shapes.add_group_shape()
        group.shapes.add_textbox(Inches(1), Inches(3), Inches(2), Inches(1)).text = 'TEXTO AGRUPADO'
        source = self.sources/'presentacion.pptx'
        presentation.save(source)
        result = convert_native(self.root, source)
        text = (Path(result['output'])/'documento.md').read_text(encoding='utf-8')
        self.assertIn('TABLA&#124;PPT', text)
        self.assertIn('TEXTO AGRUPADO', text)

    def test_batch_exports_markdown_report(self):
        self.source('a.txt', 'Texto')
        result = self.plan()
        self.assertTrue(Path(result['report']).is_file())
        result = run_batch(self.root)
        text = Path(result['report']).read_text(encoding='utf-8')
        self.assertIn('Abrir Markdown', text)
        self.assertIn('terminado', text)

    def test_batch_fault_appears_in_health(self):
        from conversion.diagnostics import run_diagnostics
        source = self.source('a.txt', 'Texto')
        self.plan()
        source.write_text('Cambió', encoding='utf-8')
        run_batch(self.root)
        checks = run_diagnostics(self.root)['checks']
        self.assertFalse(next(c for c in checks if c['name'] == 'Lotes locales')['ok'])

    def test_key_slots_are_separate_and_never_rotate_environment(self):
        with patch.dict('os.environ', {'GEMINI_API_KEY':'uno', 'GEMINI_API_KEY_3':'tres'}):
            self.assertEqual(get_key('gemini-api',3), 'tres')
            self.assertEqual(get_key('gemini-api'), 'uno')
        self.assertEqual(key_location('gemini-api', 4), ('GEMINI_API_KEY_4', 'SistemaMD/gemini-api/4'))
        self.assertEqual(key_location('deepseek-api', 5), ('DEEPSEEK_API_KEY_5', 'SistemaMD/deepseek-api/5'))
        self.assertEqual(key_location('gemini-api', 5), ('GEMINI_API_KEY_5', 'SistemaMD/gemini-api/5'))
        for provider, slot in [('deepseek-api',6), ('gemini-api',6), ('gemini-cli',1)]:
            with self.assertRaises(ValueError):
                key_location(provider, slot)

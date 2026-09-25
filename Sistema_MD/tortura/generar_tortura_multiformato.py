"""FASE 1 · Banco de pruebas de tortura multiformato para Sistema MD.

Fabrica archivos sintéticos con patologías estructurales conocidas (celdas combinadas,
caracteres que rompen Markdown, tablas anidadas, páginas rotadas, bloques CAD anidados,
imágenes que requieren OCR y archivos corruptos). Todo se genera localmente, sin red,
con datos inventados. Escribe ``esperado.json`` con el resultado exigible por archivo:

  convertir           -> debe publicarse un paquete válido
  rechazo_controlado  -> debe rechazarse con un error explicado (nunca excepción cruda)

Uso:  python generar_tortura_multiformato.py [--salida stress_test_suite]

Dependencias: openpyxl, python-docx, python-pptx, PyMuPDF, ezdxf, Pillow (las del núcleo).
Opcional: LibreOffice (``soffice``) para derivar .xls/.doc/.ppt reales desde los OOXML.
"""
from __future__ import annotations

import argparse
import datetime as dt
import io
import json
import os
from pathlib import Path
import re
import shutil
import struct
import subprocess
import sys
import tempfile
import zipfile

HERE = Path(__file__).resolve().parent
TEXTO_CONFLICTIVO = ("Pipe | dentro · barra / y \\ inversa · llaves { } · corchetes [ ] · "
                     "asteriscos *negrita falsa* · backtick `code` · triples ``` y \"\"\" comillas")
FUENTES = ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/Library/Fonts/Arial Unicode.ttf",
           "C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/segoeui.ttf")


def fuente_ttf():
    return next((ruta for ruta in FUENTES if Path(ruta).is_file()), None)


def pil_font(size):
    from PIL import ImageFont
    ruta = fuente_ttf()
    return ImageFont.truetype(ruta, size) if ruta else ImageFont.load_default()


# ================================================================== EXCEL
def _parchear_zip(ruta: Path, cambios: dict[str, callable], extra: dict[str, bytes] | None = None):
    """Reescribe partes del ZIP (texto XML) sin tocar el resto."""
    temporal = ruta.with_suffix(ruta.suffix + ".tmp")
    with zipfile.ZipFile(ruta) as origen, zipfile.ZipFile(temporal, "w", zipfile.ZIP_DEFLATED) as destino:
        for item in origen.infolist():
            data = origen.read(item.filename)
            if item.filename in cambios:
                data = cambios[item.filename](data.decode("utf-8")).encode("utf-8")
            destino.writestr(item, data)
        for nombre, data in (extra or {}).items():
            destino.writestr(nombre, data)
    os.replace(temporal, ruta)


def _partes(ruta: Path, prefijo: str) -> list[str]:
    with zipfile.ZipFile(ruta) as archivo:
        return [n for n in archivo.namelist() if n.startswith(prefijo) and n.endswith(".xml")]


def excel(destino: Path) -> Path:
    import openpyxl
    from openpyxl.styles import Alignment
    from openpyxl.worksheet.table import Table

    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Combinadas y texto"
    # Combinadas: horizontal A1:C1, vertical A2:A5 y bloque B3:D4.
    hoja["A1"] = "Título combinado horizontal A1:C1"
    hoja.merge_cells("A1:C1")
    hoja["A2"] = "Vertical\nA2:A5"
    hoja.merge_cells("A2:A5")
    hoja["B3"] = "Bloque B3:D4"
    hoja.merge_cells("B3:D4")
    hoja["B2"] = "Salto\nLF"
    hoja["C2"] = "Salto\r\nCRLF"
    hoja["D2"] = "__CR_ESCAPADO__"         # se sustituye por _x000D_ como lo guarda Excel
    hoja["B5"] = TEXTO_CONFLICTIVO
    hoja["C5"] = "| empieza con pipe | y termina |"
    hoja["D5"] = "# no es título\n> no es cita\n- no es lista"
    hoja["E1"] = "=texto que parece fórmula"
    hoja["E1"].data_type = "s"
    hoja["E2"] = "__CTRL__"                # carácter de control escapado (_x0007_)
    for celda in ("B2", "C2", "A2"):
        hoja[celda].alignment = Alignment(wrap_text=True)

    errores = libro.create_sheet("Fórmulas y errores")
    errores.append(["Caso", "Fórmula", "Esperado"])
    casos = [("N/A", "=NA()", "#N/A"), ("VALUE", '="texto"+1', "#VALUE!"),
             ("DIV/0", "=1/0", "#DIV/0!"), ("REF", "=#REF!+1", "#REF!"),
             ("Anidada", '=IF(ISERROR(VLOOKUP("x",A1:B2,3,FALSE)),"sin | dato",SUM(A1:A3))', "sin | dato"),
             ("Matriz", "=SUMPRODUCT((B2:B5>0)*(C2:C5))", "0"),
             ("Otra hoja", "='Combinadas y texto'!A1", "Título combinado horizontal A1:C1"),
             ("Texto con comillas", '=CONCATENATE("a""b","|","c")', 'a"b|c')]
    for fila, (nombre, formula, esperado) in enumerate(casos, 2):
        errores.cell(fila, 1, nombre)
        errores.cell(fila, 2, formula)
        errores.cell(fila, 3, esperado)

    fechas = libro.create_sheet("Fechas & números (2024) #1")
    fechas.append(["Descripción", "Valor"])
    formatos = [("Fecha local", dt.date(2024, 2, 29), "dd/mm/yyyy"),
                ("Fecha larga es-PE", dt.datetime(2024, 12, 31, 23, 59), '[$-es-PE]dddd, d "de" mmmm "de" yyyy hh:mm'),
                ("Hora", dt.time(7, 5, 3), "hh:mm:ss"),
                ("Científico grande", 6.02214076e23, "0.00E+00"),
                ("Científico pequeño", 1.602e-19, "0.000E+00"),
                ("Moneda", 1234567.891, '"S/ "#,##0.00'),
                ("Porcentaje", 0.0725, "0.00%"),
                ("Negativo entre paréntesis", -42.5, "#,##0.00;(#,##0.00)"),
                ("Entero muy largo", 12345678901234567890, "0")]
    for fila, (nombre, valor, formato) in enumerate(formatos, 2):
        fechas.cell(fila, 1, nombre)
        celda = fechas.cell(fila, 2, valor)
        celda.number_format = formato

    ancha = libro.create_sheet("Tabla 35 columnas")
    encabezados = [f"Col {i} | x" if i % 7 == 0 else f"Campo {i}" for i in range(1, 36)]
    encabezados[0] = "Código"
    encabezados[1] = "Descripción"
    ancha.append(encabezados)
    for fila in range(2, 9):
        valores = []
        for columna in range(1, 36):
            if (fila + columna) % 4 == 0:
                valores.append(None)            # celdas vacías intercaladas
            elif columna == 1:
                valores.append(f"P-{fila:03d}")
            else:
                valores.append(f"v{fila}.{columna}" if columna % 5 else fila * columna * 1.5)
        ancha.append(valores)
    ancha.add_table(Table(displayName="Ancha35", ref="A1:AI8"))

    oculta = libro.create_sheet("Cálculos ocultos")
    oculta.sheet_state = "hidden"
    oculta["A1"], oculta["B1"], oculta["B2"] = "Factor", 3, "=B1*2"
    muy_oculta = libro.create_sheet("Muy oculta")
    muy_oculta.sheet_state = "veryHidden"
    muy_oculta["A1"] = "Solo visible por VBA"
    libro.create_sheet("Hoja vacía")

    ruta = destino / "stress_excel.xlsx"
    libro.save(ruta)

    def cache_de_errores(xml: str) -> str:
        # openpyxl no escribe resultados cacheados: se añaden como los dejaría Excel.
        mapa = {f"B{fila}": esperado for fila, (_n, _f, esperado) in enumerate(casos, 2)}

        def poner(match):
            ref, cuerpo = match.group(1), match.group(2)
            valor = mapa.get(ref)
            if valor is None:
                return match.group(0)
            tipo = "e" if valor.startswith("#") else "str"
            cuerpo = re.sub(r"<v\s*/>|<v>.*?</v>", "", cuerpo)
            return f'<c r="{ref}" t="{tipo}">{cuerpo}<v>{valor}</v></c>'
        return re.sub(r'<c r="(B\d+)"[^>]*>(<f>.*?</f>.*?)</c>', poner, xml)

    def escapes_excel(xml: str) -> str:
        return xml.replace("__CR_ESCAPADO__", "Retorno_x000D_escapado").replace("__CTRL__", "Campana_x0007_control_x0000_nulo")

    with zipfile.ZipFile(ruta) as archivo:
        hojas = {n for n in archivo.namelist() if n.startswith("xl/worksheets/sheet")}
        compartidas = "xl/sharedStrings.xml" in archivo.namelist()
    cambios = {"xl/worksheets/sheet2.xml": cache_de_errores}
    objetivo = "xl/sharedStrings.xml" if compartidas else "xl/worksheets/sheet1.xml"
    cambios[objetivo] = escapes_excel
    _parchear_zip(ruta, cambios)
    assert hojas
    return ruta


def excel_xlsm(destino: Path, xlsx: Path) -> Path:
    """Mismo libro como XLSM: tipo de contenido macro y vbaProject.bin inerte."""
    ruta = destino / "stress_excel.xlsm"
    shutil.copy2(xlsx, ruta)

    def tipo_macro(xml):
        xml = xml.replace("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                          "application/vnd.ms-excel.sheet.macroEnabled.main+xml")
        return xml.replace("</Types>", '<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>')

    def relacion(xml):
        return xml.replace("</Relationships>", '<Relationship Id="rIdVBA" '
                           'Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" '
                           'Target="vbaProject.bin"/></Relationships>')
    _parchear_zip(ruta, {"[Content_Types].xml": tipo_macro, "xl/_rels/workbook.xml.rels": relacion},
                  {"xl/vbaProject.bin": b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 504 + b"MACRO INACTIVA"})
    return ruta


# --- XLSB mínimo (BIFF12) escrito a mano: no hay escritor libre de XLSB.
def _varint(value: int) -> bytes:
    out = bytearray()
    while True:
        byte = value & 0x7F
        value >>= 7
        out.append(byte | (0x80 if value else 0))
        if not value:
            return bytes(out)


def _rec(tipo: int, data: bytes = b"") -> bytes:
    """Registro BIFF12: tipo y tamaño en 7 bits por byte; ``tipo`` es el número de [MS-XLSB]."""
    return _varint(tipo) + _varint(len(data)) + data


def _wstr(texto: str) -> bytes:
    raw = texto.encode("utf-16-le")
    return struct.pack("<I", len(raw) // 2) + raw


def excel_xlsb(destino: Path) -> Path:
    """Libro XLSB con dos hojas (una oculta), texto conflictivo, números y una fórmula =1+2."""
    def celda_st(col, texto):
        return _rec(6, struct.pack("<II", col, 0) + _wstr(texto))

    def celda_num(col, valor):
        return _rec(5, struct.pack("<IId", col, 0, valor))

    def fila(rw):
        return _rec(0, struct.pack("<IIHHBI", rw, 0, 300, 0, 0, 0))

    def formula_num(col, valor):
        rgce = bytes([0x1E]) + struct.pack("<H", 1) + bytes([0x1E]) + struct.pack("<H", 2) + bytes([0x03])
        return _rec(9, struct.pack("<IId", col, 0, valor) + struct.pack("<H", 0)
                    + struct.pack("<I", len(rgce)) + rgce + struct.pack("<I", 0))

    def hoja(filas):
        ultima_fila = max(rw for rw, _ in filas)
        ultima_col = max(len(celdas) for _, celdas in filas) - 1
        dimension = _rec(148, struct.pack("<IIII", 0, ultima_fila, 0, ultima_col))    # BrtWsDim
        cuerpo = _rec(129) + dimension + _rec(145)          # BrtBeginSheet, BrtBeginSheetData
        for rw, celdas in filas:
            cuerpo += fila(rw) + b"".join(celdas)
        return cuerpo + _rec(146) + _rec(130)

    hoja1 = hoja([(0, [celda_st(0, "Código"), celda_st(1, "Descripción | con pipe"), celda_st(2, "Total")]),
                  (1, [celda_st(0, "P-1"), celda_st(1, "Cemento\nsaco 42,5 kg"), celda_num(2, 1.602e-19)]),
                  (2, [celda_st(0, "P-2"), celda_st(1, "`backtick` *asterisco*"), formula_num(2, 3.0)])])
    hoja2 = hoja([(0, [celda_st(0, "Oculta"), celda_num(1, 42.0)])])
    libro = (_rec(131) + _rec(143)                 # BrtBeginBook, BrtBeginBundleShs
             + _rec(156, struct.pack("<II", 0, 1) + _wstr("rId1") + _wstr("Datos XLSB"))
             + _rec(156, struct.pack("<II", 1, 2) + _wstr("rId2") + _wstr("Oculta XLSB"))
             + _rec(144) + _rec(132))
    estilos = _rec(278) + _rec(279)                  # BrtBegin/EndStyleSheet
    tipos = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
             '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
             '<Default Extension="bin" ContentType="application/vnd.ms-excel.sheet.binary.macroEnabled.main"/>'
             '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
             '<Default Extension="xml" ContentType="application/xml"/>'
             '<Override PartName="/xl/worksheets/sheet1.bin" ContentType="application/vnd.ms-excel.worksheet"/>'
             '<Override PartName="/xl/worksheets/sheet2.bin" ContentType="application/vnd.ms-excel.worksheet"/>'
             '<Override PartName="/xl/styles.bin" ContentType="application/vnd.ms-excel.styles"/></Types>')
    raiz = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
            'Target="xl/workbook.bin"/></Relationships>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.bin"/>'
            '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet2.bin"/>'
            '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.bin"/>'
            '</Relationships>')
    ruta = destino / "stress_excel.xlsb"
    with zipfile.ZipFile(ruta, "w", zipfile.ZIP_DEFLATED) as archivo:
        archivo.writestr("[Content_Types].xml", tipos)
        archivo.writestr("_rels/.rels", raiz)
        archivo.writestr("xl/workbook.bin", libro)
        archivo.writestr("xl/_rels/workbook.bin.rels", rels)
        archivo.writestr("xl/worksheets/sheet1.bin", hoja1)
        archivo.writestr("xl/worksheets/sheet2.bin", hoja2)
        archivo.writestr("xl/styles.bin", estilos)
    return ruta


# ================================================================== WORD
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _notas(documento, tipo: str, textos: list[str]):
    """Crea footnotes.xml / endnotes.xml (python-docx no tiene API para notas)."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from xml.sax.saxutils import escape
    etiqueta = "footnote" if tipo == "footnotes" else "endnote"
    cuerpo = "".join(f'<w:{etiqueta} w:type="{t}" w:id="{i}"><w:p><w:r><w:{s}/></w:r></w:p></w:{etiqueta}>'
                     for i, t, s in ((-1, "separator", "separator"), (0, "continuationSeparator", "continuationSeparator")))
    for numero, texto in enumerate(textos, 1):
        cuerpo += (f'<w:{etiqueta} w:id="{numero}"><w:p><w:r><w:t xml:space="preserve">{escape(texto)}</w:t>'
                   f'</w:r></w:p></w:{etiqueta}>')
    xml = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:{tipo} xmlns:w="{W}">{cuerpo}</w:{tipo}>'
    content_type = f"application/vnd.openxmlformats-officedocument.wordprocessingml.{tipo}+xml"
    parte = Part(PackURI(f"/word/{tipo}.xml"), content_type, xml.encode("utf-8"), documento.part.package)
    documento.part.relate_to(parte, RT.FOOTNOTES if tipo == "footnotes" else RT.ENDNOTES)


def _referencia_nota(parrafo, tipo: str, numero: int):
    from docx.oxml import parse_xml
    etiqueta = "footnoteReference" if tipo == "footnotes" else "endnoteReference"
    parrafo._p.append(parse_xml(f'<w:r xmlns:w="{W}"><w:rPr><w:vertAlign w:val="superscript"/></w:rPr>'
                                f'<w:{etiqueta} w:id="{numero}"/></w:r>'))


def _numeracion_6_niveles(documento) -> int:
    from docx.oxml import parse_xml
    formatos = ["bullet", "decimal", "lowerLetter", "bullet", "upperRoman", "lowerRoman"]
    textos = ["•", "%2.", "%3)", "▪", "%5.", "%6."]
    niveles = "".join(
        f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="{f}"/><w:lvlText w:val="{t}"/>'
        f'<w:pPr><w:ind w:left="{720 * (i + 1)}" w:hanging="360"/></w:pPr></w:lvl>'
        for i, (f, t) in enumerate(zip(formatos, textos)))
    numbering = documento.part.numbering_part.element
    numbering.insert(0, parse_xml(f'<w:abstractNum xmlns:w="{W}" w:abstractNumId="90">'
                                  f'<w:multiLevelType w:val="multilevel"/>{niveles}</w:abstractNum>'))
    numbering.append(parse_xml(f'<w:num xmlns:w="{W}" w:numId="90"><w:abstractNumId w:val="90"/></w:num>'))
    return 90


def _item_lista(documento, texto: str, nivel: int, num_id: int):
    from docx.oxml import parse_xml
    parrafo = documento.add_paragraph(texto)
    parrafo._p.get_or_add_pPr().append(parse_xml(
        f'<w:numPr xmlns:w="{W}"><w:ilvl w:val="{nivel}"/><w:numId w:val="{num_id}"/></w:numPr>'))


def _bordes_invisibles(tabla):
    from docx.oxml import parse_xml
    bordes = "".join(f'<w:{b} w:val="nil"/>' for b in ("top", "left", "bottom", "right", "insideH", "insideV"))
    tabla._tbl.tblPr.append(parse_xml(f'<w:tblBorders xmlns:w="{W}">{bordes}</w:tblBorders>'))


def word(destino: Path) -> Path:
    import docx
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import parse_xml
    from docx.shared import Pt

    documento = docx.Document()
    documento.add_heading("Documento de tortura | Word", level=1)
    _notas(documento, "footnotes", ["Nota al pie con | pipe, *asterisco* y salto\nde línea.",
                                    "Segunda nota: comillas “tipográficas” — y ```triple```."])
    _notas(documento, "endnotes", ["Nota final: fuente [1] {llaves}."])

    parrafo = documento.add_paragraph("Párrafo con nota al pie")
    _referencia_nota(parrafo, "footnotes", 1)
    parrafo.add_run(" y otra")
    _referencia_nota(parrafo, "footnotes", 2)
    parrafo.add_run(" y una nota final")
    _referencia_nota(parrafo, "endnotes", 1)

    # Formato extremo en una misma línea, con enlaces rotos.
    mezcla = documento.add_paragraph()
    mezcla.add_run("negrita ").bold = True
    mezcla.add_run("cursiva ").italic = True
    tachado = mezcla.add_run("tachado ")
    tachado.font.strike = True
    todo = mezcla.add_run("todo junto ")
    todo.bold = todo.italic = True
    todo.font.strike = True
    rid = documento.part.relate_to("https://ejemplo.invalid/ruta rota?x=|", docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK,
                                   is_external=True)
    mezcla._p.append(parse_xml(f'<w:hyperlink xmlns:w="{W}" xmlns:r="{R}" r:id="{rid}"><w:r><w:t>enlace externo roto</w:t></w:r></w:hyperlink>'))
    mezcla._p.append(parse_xml(f'<w:hyperlink xmlns:w="{W}" xmlns:r="{R}" r:id="rId9999"><w:r><w:t> enlace sin relación</w:t></w:r></w:hyperlink>'))
    mezcla._p.append(parse_xml(f'<w:hyperlink xmlns:w="{W}" w:anchor="marcador_inexistente"><w:r><w:t> ancla inexistente</w:t></w:r></w:hyperlink>'))

    documento.add_paragraph(TEXTO_CONFLICTIVO)
    for trampa in ("| a | b | c |", "# Esto no es un título", "> Esto no es una cita", "```", "---",
                   "1. Esto no es lista numerada", "<script>alert('x')</script> & entidades &amp;"):
        documento.add_paragraph(trampa)

    num_id = _numeracion_6_niveles(documento)
    for nivel in range(6):
        _item_lista(documento, f"Nivel {nivel + 1} de lista mixta — elemento con | pipe", nivel, num_id)
    for nivel in (5, 0, 3, 1):
        _item_lista(documento, f"Salto desordenado al nivel {nivel + 1}", nivel, num_id)

    estilo = documento.styles.add_style("Codigo preformateado", WD_STYLE_TYPE.PARAGRAPH)
    estilo.font.name = "Courier New"
    estilo.font.size = Pt(9)
    for linea in ("def total(filas):", "    return sum(f['monto'] for f in filas)  # | pipe",
                  "print(`backticks` * 3)", "```python"):
        documento.add_paragraph(linea, style=estilo)

    formula = documento.add_paragraph("Fórmula OMML: ")
    formula._p.append(parse_xml(
        f'<m:oMath xmlns:m="{M}" xmlns:w="{W}"><m:sSup><m:e><m:r><m:t>E</m:t></m:r></m:e>'
        f'<m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup><m:r><m:t>=</m:t></m:r>'
        f'<m:f><m:num><m:r><m:t>a|b</m:t></m:r></m:num><m:den><m:r><m:t>c</m:t></m:r></m:den></m:f></m:oMath>'))

    # Tabla con combinaciones (gridSpan y vMerge), bordes invisibles y tabla anidada.
    tabla = documento.add_table(rows=4, cols=4)
    _bordes_invisibles(tabla)
    for fila in range(4):
        for columna in range(4):
            tabla.cell(fila, columna).text = f"F{fila}C{columna}"
    tabla.cell(0, 0).merge(tabla.cell(0, 2)).text = "gridSpan 3 | con pipe"
    tabla.cell(1, 0).merge(tabla.cell(3, 0)).text = "vMerge\nvertical"
    tabla.cell(1, 1).text = "Línea 1\nLínea 2\r\nLínea 3"
    interior = tabla.cell(2, 2).add_table(rows=2, cols=2)
    for fila in range(2):
        for columna in range(2):
            interior.cell(fila, columna).text = f"anidada {fila}.{columna} |"
    profunda = interior.cell(1, 1).add_table(rows=1, cols=2)
    profunda.cell(0, 0).text = "nivel 3"
    profunda.cell(0, 1).text = "`profunda`"

    # Fila con gridBefore: menos celdas físicas que columnas de la rejilla.
    irregular = documento.add_table(rows=2, cols=3)
    irregular.cell(0, 0).text, irregular.cell(0, 1).text, irregular.cell(0, 2).text = "A", "B", "C"
    fila = irregular.rows[1]._tr
    tr_pr = fila.get_or_add_trPr()
    tr_pr.append(parse_xml(f'<w:gridBefore xmlns:w="{W}" w:val="1"/>'))
    fila.remove(fila.tc_lst[0])
    irregular.rows[1].cells[-1].text = "fila con gridBefore"

    documento.add_page_break()
    documento.add_paragraph("")                  # párrafo vacío tras salto de página
    ruta = destino / "stress_word.docx"
    documento.save(ruta)
    return ruta


# ================================================================== POWERPOINT
def powerpoint(destino: Path) -> Path:
    from PIL import Image
    from pptx import Presentation
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.util import Inches, Pt

    presentacion = Presentation()
    blanco = presentacion.slide_layouts[6]

    # 1. Formas superpuestas y "SmartArt" simulado con grupos anidados.
    diapositiva = presentacion.slides.add_slide(presentacion.slide_layouts[5])
    diapositiva.shapes.title.text = "Formas superpuestas | grupos"
    grupo = diapositiva.shapes.add_group_shape()
    for indice in range(4):
        forma = grupo.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1 + indice * 0.4), Inches(2 + indice * 0.3),
                                       Inches(3), Inches(1))
        forma.text_frame.text = f"Paso {indice + 1} | superpuesto *{indice}*"
    subgrupo = grupo.shapes.add_group_shape()
    subgrupo.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5), Inches(4), Inches(1), Inches(1)).text_frame.text = "# nodo anidado"
    diapositiva.shapes.add_textbox(Inches(1), Inches(6), Inches(8), Inches(1)).text_frame.text = "| a | b |\n```\n> cita"

    # 2. Tabla con celdas combinadas y texto rotado.
    diapositiva = presentacion.slides.add_slide(presentacion.slide_layouts[5])
    diapositiva.shapes.title.text = "Tabla combinada y rotada"
    tabla = diapositiva.shapes.add_table(4, 5, Inches(0.5), Inches(1.5), Inches(9), Inches(3)).table
    for fila in range(4):
        for columna in range(5):
            tabla.cell(fila, columna).text = f"r{fila}c{columna}"
    tabla.cell(0, 0).merge(tabla.cell(0, 2))
    tabla.cell(0, 0).text = "Combinada 3 columnas | pipe"
    tabla.cell(1, 4).merge(tabla.cell(3, 4))
    tabla.cell(1, 4).text = "Vertical\nrotado"
    tabla.cell(1, 4)._tc.get_or_add_tcPr().set("vert", "vert270")
    tabla.cell(2, 1).text = "Salto\nde línea\r\ny CRLF"

    # 3. Solo imágenes y conectores sin texto nativo.
    diapositiva = presentacion.slides.add_slide(blanco)
    imagen = io.BytesIO()
    Image.new("RGB", (60, 40), (30, 120, 200)).save(imagen, format="PNG")
    imagen.seek(0)
    diapositiva.shapes.add_picture(imagen, Inches(1), Inches(1), Inches(2))
    a = diapositiva.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4), Inches(1), Inches(1), Inches(1))
    b = diapositiva.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(4), Inches(4), Inches(1), Inches(1))
    conector = diapositiva.shapes.add_connector(MSO_CONNECTOR.ELBOW, 0, 0, 0, 0)
    conector.begin_connect(a, 2)
    conector.end_connect(b, 0)
    diapositiva.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(6), Inches(1), Inches(8), Inches(3))

    # 4. Notas del orador extensas con caracteres especiales.
    diapositiva = presentacion.slides.add_slide(presentacion.slide_layouts[1])
    diapositiva.shapes.title.text = "Notas extensas"
    diapositiva.placeholders[1].text = "Viñeta 1\nViñeta 2 | pipe\n\tViñeta con tabulador"
    notas = "\n".join(f"Línea {i}: {TEXTO_CONFLICTIVO} ñ ü 😀 <etiqueta> &amp;" for i in range(1, 25))
    diapositiva.notes_slide.notes_text_frame.text = notas + "\n\n\n| tabla | falsa |\n# título falso"

    # 5. Diapositiva vacía y 6. diapositiva oculta.
    presentacion.slides.add_slide(blanco)
    oculta = presentacion.slides.add_slide(presentacion.slide_layouts[5])
    oculta.shapes.title.text = "Diapositiva oculta"
    oculta._element.set("show", "0")
    caja = oculta.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    caja.text_frame.text = "Contenido oculto con `backticks`"
    caja.text_frame.paragraphs[0].runs[0].font.size = Pt(28)

    ruta = destino / "stress_powerpoint.pptx"
    presentacion.save(ruta)
    return ruta


# ================================================================== PDF
def pdf(destino: Path) -> Path:
    import fitz
    from PIL import Image, ImageDraw

    fuente = fuente_ttf()
    documento = fitz.open()
    opciones = {"fontsize": 9, "fontname": "dejavu", "fontfile": fuente} if fuente else {"fontsize": 9}
    lorem = ("Texto de columna con ligaduras ﬁnal y ﬂujo, comillas “tipográficas” y ‘simples’ — "
             "raya em — y guion – corto. ") * 6

    # Página 1: dos columnas arriba, tres abajo.
    pagina = documento.new_page(width=595, height=842)
    pagina.insert_textbox(fitz.Rect(40, 40, 555, 70), "Página a dos y tres columnas", **{**opciones, "fontsize": 14})
    for i, x in enumerate((40, 310)):
        pagina.insert_textbox(fitz.Rect(x, 80, x + 245, 420), f"COLUMNA {i + 1} DE 2. " + lorem, **opciones)
    for i, x in enumerate((40, 220, 400)):
        pagina.insert_textbox(fitz.Rect(x, 440, x + 160, 800), f"COLUMNA {i + 1} DE 3. " + lorem, **opciones)

    # Página 2: tabla sin bordes pegada a un párrafo.
    pagina = documento.new_page(width=595, height=842)
    pagina.insert_textbox(fitz.Rect(40, 40, 555, 120), "Párrafo inmediatamente anterior a una tabla sin bordes. "
                          "Los valores se alinean solo por posición.", **opciones)
    filas = [("Partida", "Unidad", "Cantidad", "Precio"), ("Cemento | tipo I", "bls", "1 250", "S/ 32,50"),
             ("Arena gruesa", "m³", "84,5", "S/ 60,00"), ("Acero ø 1/2\"", "kg", "3 400", "S/ 4,10")]
    for fila, valores in enumerate(filas):
        for columna, valor in enumerate(valores):
            pagina.insert_text((40 + columna * 130, 150 + fila * 18), valor, **opciones)
    pagina.insert_textbox(fitz.Rect(40, 240, 555, 320), "Párrafo inmediatamente posterior, sin separación.", **opciones)

    # Páginas 3 y 4: rotadas 90° y 270°.
    for grados in (90, 270):
        pagina = documento.new_page(width=595, height=842)
        pagina.insert_textbox(fitz.Rect(40, 40, 555, 400), f"Página rotada {grados}°. " + lorem, **opciones)
        pagina.set_rotation(grados)

    # Página 5: "escaneada" (solo imagen, sin capa de texto) que requiere OCR.
    lienzo = Image.new("L", (1240, 1754), 255)
    dibujo = ImageDraw.Draw(lienzo)
    letra = pil_font(44)
    for linea, texto in enumerate(("INFORME ESCANEADO", "Cemento Portland tipo I", "Cantidad 1250 bolsas",
                                   "Precio unitario 32,50")):
        dibujo.text((120, 160 + linea * 110), texto, fill=0, font=letra)
    png = io.BytesIO()
    lienzo.save(png, format="PNG")
    pagina = documento.new_page(width=595, height=842)
    pagina.insert_image(pagina.rect, stream=png.getvalue())

    # Página 6: caracteres especiales y control embebido en el flujo de texto.
    pagina = documento.new_page(width=595, height=842)
    pagina.insert_textbox(fitz.Rect(40, 40, 555, 800),
                          "Ligaduras: ﬁ ﬂ ﬀ ﬃ · Comillas: “ ” ‘ ’ « » · Rayas: — – · Símbolos: ± × ÷ ≤ ≥ ∑ √ ∞ · "
                          + TEXTO_CONFLICTIVO, **opciones)
    ruta = destino / "stress_pdf.pdf"
    documento.save(ruta)
    documento.close()
    return ruta


# ================================================================== CAD
def cad(destino: Path) -> Path:
    import ezdxf

    dibujo = ezdxf.new("R2018", setup=True)
    capas = dibujo.layers
    capas.add("CONGELADA", color=1).freeze()
    capas.add("APAGADA", color=2).off()
    capas.add("BLOQUEADA", color=3).lock()
    capas.add("COTAS", color=4)
    modelo = dibujo.modelspace()

    interior = dibujo.blocks.new("PERNO_INTERIOR")
    interior.add_circle((0, 0), 0.5)
    interior.add_text("perno | M12", dxfattribs={"height": 0.2}).set_placement((0.6, 0))
    exterior = dibujo.blocks.new("PLACA_EXTERIOR")
    exterior.add_lwpolyline([(0, 0), (4, 0), (4, 3), (0, 3)], close=True)
    for x, y in ((0.5, 0.5), (3.5, 0.5), (3.5, 2.5), (0.5, 2.5)):
        exterior.add_blockref("PERNO_INTERIOR", (x, y))           # INSERT dentro de INSERT
    exterior.add_attdef("CODIGO", (0, -0.5), dxfattribs={"height": 0.25})
    for indice, (x, rot, escala) in enumerate(((0, 0, 1), (10, 30, 2), (25, 90, 0.5))):
        referencia = modelo.add_blockref("PLACA_EXTERIOR", (x, 0), dxfattribs={"rotation": rot,
                                         "xscale": escala, "yscale": escala})
        referencia.add_auto_attribs({"CODIGO": f"PL-{indice + 1} | *placa*"})

    modelo.add_text("Texto en capa congelada", dxfattribs={"layer": "CONGELADA", "height": 1}).set_placement((0, 20))
    modelo.add_text("Texto en capa apagada", dxfattribs={"layer": "APAGADA", "height": 1}).set_placement((0, 22))
    modelo.add_text("Texto en capa bloqueada", dxfattribs={"layer": "BLOQUEADA", "height": 1}).set_placement((0, 24))
    modelo.add_mtext("MTEXT línea 1\\PLínea 2 con \\fArial|b1|i0|c0|p34;negrita\\f;\\P"
                     "\\A1;centrado vertical {\\C1;rojo} \\Ssup^sub; 100% | pipe \\~espacio duro",
                     dxfattribs={"char_height": 0.5}).set_location((0, 30))
    cota = modelo.add_linear_dim(base=(0, -5), p1=(0, 0), p2=(40, 0), dimstyle="EZDXF",
                                 override={"dimtxt": 0.5}, dxfattribs={"layer": "COTAS"})
    cota.set_text("<> mm | cota")
    cota.render()
    modelo.add_aligned_dim(p1=(0, 0), p2=(10, 10), distance=2, dxfattribs={"layer": "COTAS"}).render()
    modelo.add_lwpolyline([(0, 40, 0, 0, 0.5), (10, 40, 0, 0, -1.0), (20, 45, 0, 0, 0.0), (20, 55, 0, 0, 0.414)],
                          format="xyseb", close=True)
    modelo.add_polyline3d([(0, 60, 0), (5, 62, 3), (10, 60, 6), (15, 65, -2)])
    # POLYLINE 2D (entidad clásica, no LW) con arcos por vértice y cerrada. En DXF la 3D no
    # admite bulge: sus tramos son siempre rectos.
    modelo.add_polyline2d([(30, 40, 0, 0, 0.6), (35, 45, 0, 0, -0.4), (40, 40, 0, 0, 1.0)],
                          format="xyseb").close(True)
    presentacion = dibujo.layouts.new("Lámina A1 | revisión")
    presentacion.add_text("Rótulo en paper space", dxfattribs={"height": 2}).set_placement((10, 10))
    ruta = destino / "stress_cad.dxf"
    dibujo.saveas(ruta)
    return ruta


# ================================================================== IMÁGENES
def imagenes(destino: Path) -> list[Path]:
    from PIL import Image, ImageDraw

    def lamina(texto, tamano=(1400, 500)):
        lienzo = Image.new("RGB", tamano, "white")
        dibujo = ImageDraw.Draw(lienzo)
        for linea, fila in enumerate(texto.splitlines()):
            dibujo.text((60, 60 + linea * 90), fila, fill="black", font=pil_font(56))
        return lienzo

    rutas = []
    lamina("PLANO DE OBRA 2024\nCemento 1250 bolsas\nAcero 3400 kg").save(destino / "stress_ocr.png")
    rutas.append(destino / "stress_ocr.png")
    lamina("FACTURA 000123\nTotal S/ 4.100,50").save(destino / "stress_ocr.jpg", quality=85)
    rutas.append(destino / "stress_ocr.jpg")
    paginas = [lamina("TIFF PAGINA 1\nPresupuesto"), lamina("TIFF PAGINA 2\nMetrado final")]
    paginas[0].save(destino / "stress_ocr.tiff", save_all=True, append_images=paginas[1:], compression="tiff_lzw")
    rutas.append(destino / "stress_ocr.tiff")
    Image.new("RGB", (300, 200), "white").save(destino / "stress_imagen_en_blanco.png")
    rutas.append(destino / "stress_imagen_en_blanco.png")
    return rutas


# ================================================================== CORRUPTOS
def corruptos(destino: Path, xlsx: Path, docx: Path, pdf_ruta: Path, png: Path) -> dict[str, str]:
    carpeta = destino
    datos = xlsx.read_bytes()
    (carpeta / "corrupto_zip_truncado.xlsx").write_bytes(datos[: len(datos) // 2])
    (carpeta / "corrupto_vacio.docx").write_bytes(b"")
    (carpeta / "corrupto_basura.pdf").write_bytes(b"%PDF-1.7\n" + bytes(range(256)) * 8)
    png_datos = png.read_bytes()
    (carpeta / "corrupto_png_truncado.png").write_bytes(png_datos[: len(png_datos) // 3])
    with zipfile.ZipFile(docx) as origen, zipfile.ZipFile(carpeta / "corrupto_xml_roto.docx", "w") as copia:
        for item in origen.infolist():
            contenido = origen.read(item.filename)
            if item.filename == "word/document.xml":
                contenido = contenido[: len(contenido) // 2]
            copia.writestr(item, contenido)
    shutil.copy2(xlsx, carpeta / "mentira_excel_llamado.pptx")          # la extensión miente
    (carpeta / "stress_cad_falso.dwg").write_bytes(b"AC1032" + b"\x00\x13\x37" * 400)
    dxf = (destino / "stress_cad.dxf").read_bytes()
    (carpeta / "corrupto_dxf_truncado.dxf").write_bytes(dxf[: len(dxf) // 2])
    return {"corrupto_zip_truncado.xlsx": "rechazo_controlado", "corrupto_vacio.docx": "rechazo_controlado",
            "corrupto_basura.pdf": "rechazo_controlado", "corrupto_png_truncado.png": "rechazo_controlado",
            "corrupto_xml_roto.docx": "rechazo_controlado", "mentira_excel_llamado.pptx": "convertir",
            "stress_cad_falso.dwg": "rechazo_controlado", "corrupto_dxf_truncado.dxf": "rechazo_controlado"}


# ================================================================== LEGADO vía LibreOffice
def legado(destino: Path, fuentes: dict[str, Path]) -> dict[str, Path]:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:                        # Windows: LibreOffice no suele estar en el PATH
        candidata = Path(os.environ.get("ProgramFiles", "C:/Program Files")) / "LibreOffice/program/soffice.exe"
        soffice = str(candidata) if candidata.is_file() else None
    if not soffice:
        print("[aviso] LibreOffice no disponible: no se generan .xls/.doc/.ppt", file=sys.stderr)
        return {}
    creados = {}
    with tempfile.TemporaryDirectory(prefix="tortura_lo_") as perfil:
        for extension, origen in fuentes.items():
            subprocess.run([soffice, "--headless", f"-env:UserInstallation={Path(perfil).as_uri()}",
                            "--convert-to", extension, "--outdir", str(destino), str(origen)],
                           check=False, capture_output=True, timeout=300)
            resultado = destino / (origen.stem + "." + extension)
            if resultado.is_file():
                creados[resultado.name] = resultado
            else:
                print(f"[aviso] LibreOffice no produjo {resultado.name}", file=sys.stderr)
    return creados


# Canales que no pueden desaparecer en silencio (regla F del validador).
EXCEL_OOXML = ["Título combinado horizontal A1:C1", "Vertical A2:A5", "Salto CRLF", "Retorno escapado", "Campana",
               "Pipe | dentro", "llaves { }", "corchetes [ ]", "triples ``` y", "#N/A", "#DIV/0!", "#REF!",
               "Cálculos ocultos", "Solo visible por VBA", "6.02214076e+23", "2024-02-29", "Campo 34", "Col 35 | x"]
MARCADORES = {
    "stress_excel.xlsx": EXCEL_OOXML,
    "stress_excel.xlsm": EXCEL_OOXML,
    "mentira_excel_llamado.pptx": ["Título combinado horizontal A1:C1", "Solo visible por VBA"],
    "stress_excel.xls": ["Título combinado horizontal A1:C1", "Pipe | dentro", "Cálculos ocultos", "Campo 34"],
    "stress_excel.xlsb": ["Datos XLSB", "Oculta XLSB", "Descripción | con pipe", "Cemento saco 42,5 kg", "1.602e-19"],
    "stress_word.docx": ["Nota al pie con | pipe", "Segunda nota", "Nota final: fuente [1] {llaves}", "anidada 0.0 |",
                         "nivel 3", "`profunda`", "fila con gridBefore", "Nivel 6 de lista mixta", "a|b",
                         "enlace externo roto", "enlace sin relación", "ancla inexistente", "def total(filas):",
                         "| a | b | c |", "# Esto no es un título", "gridSpan 3 | con pipe"],
    "stress_word.doc": ["Nota al pie con | pipe", "anidada 0.0 |", "Nivel 6 de lista mixta", "gridSpan 3 | con pipe"],
    "stress_powerpoint.pptx": ["Paso 4 | superpuesto", "# nodo anidado", "Combinada 3 columnas | pipe", "Vertical rotado",
                               "Viñeta 2 | pipe", "Línea 24:", "título falso", "Diapositiva oculta",
                               "Contenido oculto con `backticks`", "Diapositiva 6 *(oculta)*", "| a | b |"],
    "stress_powerpoint.ppt": ["Paso 4 | superpuesto", "Combinada 3 columnas | pipe", "Línea 24:", "Diapositiva oculta"],
    "stress_pdf.pdf": ["COLUMNA 1 DE 2", "COLUMNA 3 DE 3", "Cemento | tipo I", "Página rotada 90°", "Página rotada 270°",
                       "INFORME ESCANEADO", "“tipográficas”", "raya em —"],
    "stress_cad.dxf": ["perno | M12", "PL-1 | *placa*", "Texto en capa congelada", "Texto en capa apagada",
                       "Texto en capa bloqueada", "MTEXT línea 1", "Línea 2 con", "negrita", "Rótulo en paper space",
                       "mm | cota"],
    "stress_ocr.png": ["PLANO DE OBRA 2024", "Cemento 1250 bolsas"],
    "stress_ocr.jpg": ["FACTURA 000123"],
    "stress_ocr.tiff": ["TIFF PAGINA 1", "TIFF PAGINA 2"],
}


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--salida", type=Path, default=HERE / "stress_test_suite")
    args = parser.parse_args(argv)
    destino = args.salida.resolve()
    if destino.exists():
        shutil.rmtree(destino)
    destino.mkdir(parents=True)

    esperado: dict[str, dict] = {}

    def registrar(ruta: Path, familia: str, resultado: str = "convertir"):
        esperado[ruta.name] = {"familia": familia, "esperado": resultado,
                               "debe_contener": MARCADORES.get(ruta.name, [])}

    xlsx = excel(destino)
    registrar(xlsx, "excel")
    registrar(excel_xlsm(destino, xlsx), "excel")
    registrar(excel_xlsb(destino), "excel")
    docx = word(destino)
    registrar(docx, "word")
    pptx = powerpoint(destino)
    registrar(pptx, "powerpoint")
    pdf_ruta = pdf(destino)
    registrar(pdf_ruta, "pdf")
    registrar(cad(destino), "cad")
    fotos = imagenes(destino)
    for foto in fotos:
        registrar(foto, "imagen")
    # python-pptx guarda «\r» como «_x000D_»; LibreOffice lo lee como texto literal y el .ppt
    # heredaría ese literal. El .ppt se deriva de una copia sin ese escape: la tortura del
    # escape queda en el .pptx, donde sí tiene significado.
    pptx_para_ppt = destino / "_pptx_para_ppt" / pptx.name
    pptx_para_ppt.parent.mkdir()
    shutil.copy2(pptx, pptx_para_ppt)
    _parchear_zip(pptx_para_ppt, {n: (lambda xml: xml.replace("_x000D_", "")) for n in _partes(pptx, "ppt/slides/slide")})
    for nombre, ruta in legado(destino, {"xls": xlsx, "doc": docx, "ppt": pptx_para_ppt}).items():
        registrar(ruta, {"xls": "excel", "doc": "word", "ppt": "powerpoint"}[ruta.suffix[1:]])
    shutil.rmtree(pptx_para_ppt.parent)
    for nombre, resultado in corruptos(destino, xlsx, docx, pdf_ruta, fotos[0]).items():
        registrar(destino / nombre, "corrupto" if resultado != "convertir" else "excel", resultado)

    (destino / "esperado.json").write_text(json.dumps(esperado, ensure_ascii=False, indent=2, sort_keys=True),
                                          encoding="utf-8")
    print(json.dumps({"salida": str(destino), "archivos": len(esperado)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
